from __future__ import annotations

import csv
import json
import math
import re
import shutil
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
ASSET_DIR = OUT_DIR / "report_assets"
OUTPUT = OUT_DIR / "多无人机三维编队与协同避障课程大作业报告.docx"

FONT_CN = "宋体"
FONT_EN = "Times New Roman"
FONT_MATH = "Cambria Math"


def set_run_font(run, cn=FONT_CN, en=FONT_EN, size=12, bold=False, italic=False):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    rfonts.set(qn("w:eastAsia"), cn)
    rfonts.set(qn("w:cs"), en)
    return run


def set_east_asia_font(style, cn=FONT_CN, en=FONT_EN):
    style.font.name = en
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    rfonts.set(qn("w:eastAsia"), cn)
    rfonts.set(qn("w:cs"), en)


def set_outline_level(style, level):
    ppr = style.element.get_or_add_pPr()
    old = ppr.find(qn("w:outlineLvl"))
    if old is not None:
        ppr.remove(old)
    el = OxmlElement("w:outlineLvl")
    el.set(qn("w:val"), str(level))
    ppr.append(el)


def set_keep(style, keep_next=True, keep_lines=True):
    ppr = style.element.get_or_add_pPr()
    for tag, val in (("w:keepNext", keep_next), ("w:keepLines", keep_lines)):
        old = ppr.find(qn(tag))
        if old is not None:
            ppr.remove(old)
        if val:
            ppr.append(OxmlElement(tag))


def disable_auto_spacing_ppr(ppr):
    for tag in ("w:autoSpaceDE", "w:autoSpaceDN"):
        node = ppr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            ppr.append(node)
        node.set(qn("w:val"), "0")


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **kwargs):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        data = kwargs.get(edge)
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        if data is None:
            element.set(qn("w:val"), "nil")
        else:
            element.set(qn("w:val"), data.get("val", "single"))
            element.set(qn("w:sz"), str(data.get("sz", 8)))
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), data.get("color", "000000"))


def set_table_borders_none(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_col_width(cell, width_cm):
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def strip_alnum_han_spaces(text):
    text = re.sub(r"(?<=[A-Za-z0-9%°]) +(?=[\u3400-\u9fff])", "", text)
    text = re.sub(r"(?<=[\u3400-\u9fff]) +(?=[A-Za-z0-9])", "", text)
    return text


def add_complex_field(paragraph, instruction, result_text, bookmark=None, font_size=12):
    if bookmark:
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(add_complex_field.bookmark_id))
        start.set(qn("w:name"), bookmark)
        paragraph._p.append(start)
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    begin_run._r.append(begin)
    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    instr_run._r.append(instr)
    sep_run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    sep_run._r.append(sep)
    result_run = paragraph.add_run(result_text)
    set_run_font(result_run, size=font_size)
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)
    if bookmark:
        end_bm = OxmlElement("w:bookmarkEnd")
        end_bm.set(qn("w:id"), str(add_complex_field.bookmark_id))
        paragraph._p.append(end_bm)
        add_complex_field.bookmark_id += 1
    return result_run


add_complex_field.bookmark_id = 10


def add_ref(paragraph, bookmark, cached_number):
    add_complex_field(paragraph, f" REF {bookmark} \\h ", str(cached_number), font_size=12)


def add_seq_caption(doc, label, number, title, bookmark, before=False):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = bool(before)
    r = p.add_run(label)
    set_run_font(r, size=10.5)
    add_complex_field(p, f" SEQ {label} \\* ARABIC ", str(number), bookmark=bookmark, font_size=10.5)
    r = p.add_run(strip_alnum_han_spaces(title))
    set_run_font(r, size=10.5)
    return p


def add_toc(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    begin_run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    begin_run._r.append(begin)
    instr_run = p.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    instr_run._r.append(instr)
    sep_run = p.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    sep_run._r.append(sep)
    cached = p.add_run("目录将在打开文档时自动更新；如未更新，请按Ctrl+A后按F9。")
    set_run_font(cached, size=10.5)
    end_run = p.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def add_page_field(paragraph):
    add_complex_field(paragraph, " PAGE ", "1", font_size=9)


def add_body(doc, text, style="正文", first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    if first_line:
        p.paragraph_format.first_line_indent = Pt(24)
    r = p.add_run(strip_alnum_han_spaces(text))
    set_run_font(r, size=12)
    return p


def add_body_with_refs(doc, parts, first_line=True):
    p = doc.add_paragraph(style="正文")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_line:
        p.paragraph_format.first_line_indent = Pt(24)
    for part in parts:
        if isinstance(part, tuple) and part[0] == "ref":
            add_ref(p, part[1], part[2])
        else:
            r = p.add_run(strip_alnum_han_spaces(str(part)))
            set_run_font(r, size=12)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_equation(doc, text):
    p = doc.add_paragraph(style="公式")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, cn=FONT_CN, en=FONT_MATH, size=11.5, italic=True)
    return p


def add_three_line_table(doc, headers, rows, widths, font_size=9.5, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders_none(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_cant_split(hdr)
    for i, (cell, header, width) in enumerate(zip(hdr.cells, headers, widths)):
        set_col_width(cell, width)
        set_cell_margins(cell, top=100, start=70, bottom=100, end=70)
        set_cell_border(
            cell,
            top={"sz": 12, "color": "000000"},
            bottom={"sz": 8, "color": "000000"},
        )
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(strip_alnum_han_spaces(str(header)))
        set_run_font(r, size=font_size, bold=True)
    for row_data in rows:
        row = table.add_row()
        set_cant_split(row)
        for i, (cell, value, width) in enumerate(zip(row.cells, row_data, widths)):
            set_col_width(cell, width)
            set_cell_margins(cell, top=85, start=70, bottom=85, end=70)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            if alignments and i < len(alignments):
                p.alignment = alignments[i]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(strip_alnum_han_spaces(str(value)))
            set_run_font(r, size=font_size)
    for i, cell in enumerate(table.rows[-1].cells):
        set_cell_border(cell, bottom={"sz": 12, "color": "000000"})
    return table


def add_figure(doc, image_path, width_cm=14.5, alt=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    inline = run.add_picture(str(image_path), width=Cm(width_cm))
    doc_pr = inline._inline.docPr
    doc_pr.set("descr", alt or image_path.stem)
    return p


def add_bullets(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph(style="项目符号")
        r = p.add_run(f"（{index}）{strip_alnum_han_spaces(item)}")
        set_run_font(r, size=12)


def read_csv(path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def draw_flow_diagram(path):
    width, height = 1800, 940
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    font_path = Path("C:/Windows/Fonts/simsun.ttc")
    bold_path = Path("C:/Windows/Fonts/simhei.ttf")
    f_title = ImageFont.truetype(str(bold_path if bold_path.exists() else font_path), 52)
    f_box = ImageFont.truetype(str(font_path), 37)
    f_small = ImageFont.truetype(str(font_path), 30)
    blue = "#24527A"
    light = "#EAF2F8"
    green = "#E8F5E9"
    orange = "#FFF3E0"
    gray = "#5B6573"
    title = "多无人机编队与协同避障控制链"
    bbox = d.textbbox((0, 0), title, font=f_title)
    d.text(((width - bbox[2]) / 2, 45), title, font=f_title, fill="#111111")
    boxes = [
        (95, 220, 390, 390, "任务输入", "目标点、障碍物\n编队构型与参数", light),
        (465, 220, 760, 390, "控制量计算", "目标吸引、编队保持\n障碍排斥、机间分离", green),
        (835, 220, 1130, 390, "约束处理", "速度限幅、平滑更新\n局部极小检测与逃逸", orange),
        (1205, 220, 1500, 390, "状态更新", "三维位置、速度\n移动障碍物真实位置", light),
        (650, 585, 945, 755, "评价与输出", "安全距离、编队误差\n完成时间、轨迹与动画", green),
    ]
    for x1, y1, x2, y2, head, body, fill in boxes:
        d.rounded_rectangle((x1, y1, x2, y2), radius=24, fill=fill, outline=blue, width=4)
        hb = d.textbbox((0, 0), head, font=f_box)
        d.text(((x1 + x2 - hb[2]) / 2, y1 + 22), head, font=f_box, fill=blue)
        lines = body.split("\n")
        for idx, line in enumerate(lines):
            tb = d.textbbox((0, 0), line, font=f_small)
            d.text(((x1 + x2 - tb[2]) / 2, y1 + 86 + idx * 48), line, font=f_small, fill=gray)
    def arrow(x1, y1, x2, y2):
        d.line((x1, y1, x2, y2), fill=blue, width=7)
        ang = math.atan2(y2 - y1, x2 - x1)
        l = 22
        for a in (ang + 2.55, ang - 2.55):
            d.line((x2, y2, x2 + l * math.cos(a), y2 + l * math.sin(a)), fill=blue, width=7)
    arrow(390, 305, 465, 305)
    arrow(760, 305, 835, 305)
    arrow(1130, 305, 1205, 305)
    arrow(1350, 390, 945, 640)
    arrow(650, 670, 245, 390)
    d.text((133, 490), "轨迹状态反馈", font=f_small, fill=gray)
    img.save(path, dpi=(220, 220))


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    set_east_asia_font(normal)
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    set_east_asia_font(title)
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(20)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    h1 = styles["Heading 1"]
    set_east_asia_font(h1)
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(14)
    h1.paragraph_format.page_break_before = True
    set_outline_level(h1, 0)
    set_keep(h1)

    h2 = styles["Heading 2"]
    set_east_asia_font(h2)
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(8)
    set_outline_level(h2, 1)
    set_keep(h2)

    h3 = styles["Heading 3"]
    set_east_asia_font(h3)
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0, 0, 0)
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(6)
    set_outline_level(h3, 2)
    set_keep(h3)

    for name in ("TOC 1", "TOC 2", "TOC 3"):
        if name in styles:
            s = styles[name]
            set_east_asia_font(s)
            s.font.size = Pt(10.5)
            s.font.color.rgb = RGBColor(0, 0, 0)

    caption = styles["Caption"]
    set_east_asia_font(caption)
    caption.font.size = Pt(10.5)
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.font.italic = False
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_lines = True

    body = styles.add_style("正文", WD_STYLE_TYPE.PARAGRAPH)
    set_east_asia_font(body)
    body.font.size = Pt(12)
    body.font.color.rgb = RGBColor(0, 0, 0)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    body.paragraph_format.space_after = Pt(0)
    body.paragraph_format.widow_control = True

    equation = styles.add_style("公式", WD_STYLE_TYPE.PARAGRAPH)
    set_east_asia_font(equation, cn=FONT_CN, en=FONT_MATH)
    equation.font.size = Pt(11.5)
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.space_before = Pt(4)
    equation.paragraph_format.space_after = Pt(4)

    bullet = styles.add_style("项目符号", WD_STYLE_TYPE.PARAGRAPH)
    set_east_asia_font(bullet)
    bullet.font.size = Pt(12)
    bullet.paragraph_format.left_indent = Pt(24)
    bullet.paragraph_format.first_line_indent = Pt(-18)
    bullet.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    bullet.paragraph_format.space_after = Pt(0)

    abstract = styles.add_style("摘要正文", WD_STYLE_TYPE.PARAGRAPH)
    set_east_asia_font(abstract)
    abstract.font.size = Pt(12)
    abstract.paragraph_format.first_line_indent = Pt(24)
    abstract.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    abstract.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style in styles:
        if style.type == WD_STYLE_TYPE.PARAGRAPH:
            disable_auto_spacing_ppr(style.element.get_or_add_pPr())


def setup_sections(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)


def set_update_fields(doc):
    settings = doc.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(p)


def disable_auto_spacing_everywhere(doc):
    parts = [doc.element.body]
    for section in doc.sections:
        parts.extend([
            section.header._element,
            section.footer._element,
            section.first_page_header._element,
            section.first_page_footer._element,
            section.even_page_header._element,
            section.even_page_footer._element,
        ])
    seen = set()
    for part in parts:
        marker = id(part)
        if marker in seen:
            continue
        seen.add(marker)
        for paragraph in part.iter(qn("w:p")):
            ppr = paragraph.find(qn("w:pPr"))
            if ppr is None:
                ppr = OxmlElement("w:pPr")
                paragraph.insert(0, ppr)
            disable_auto_spacing_ppr(ppr)


def build_document():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    flow_path = ASSET_DIR / "system_flow.png"
    draw_flow_diagram(flow_path)

    exp1 = read_csv(ROOT / "results/experiment_1/metrics_summary.csv")
    exp2 = read_csv(ROOT / "results/experiment_2/metrics_summary.csv")
    exp3 = read_csv(ROOT / "results/experiment_3/summary.csv")
    exp4 = read_csv(ROOT / "results/experiment_4/summary.csv")
    val3 = json.loads((ROOT / "results/experiment_3/validation_summary.json").read_text(encoding="utf-8"))
    val4 = json.loads((ROOT / "results/experiment_4/validation_summary.json").read_text(encoding="utf-8"))

    doc = Document()
    configure_styles(doc)
    setup_sections(doc)
    set_update_fields(doc)
    cp = doc.core_properties
    cp.title = "基于人工势场法的多无人机三维编队与协同避障研究"
    cp.subject = "课程大作业报告"
    cp.author = ""
    cp.last_modified_by = ""
    cp.keywords = "多无人机；三维编队；人工势场；协同避障；局部极小；预测避障"
    cp.comments = ""

    # 封面
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.add_run("基于人工势场法的多无人机三维编队与协同避障研究")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("课程大作业报告")
    set_run_font(r, size=16, bold=True)
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("研究对象：5架无人机三维楔形编队")
    set_run_font(r, size=12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("仿真平台：Python、NumPy与Matplotlib")
    set_run_font(r, size=12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("2026年10月")
    set_run_font(r, size=12)
    doc.add_page_break()

    # 中文摘要
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("摘  要")
    set_run_font(r, size=16, bold=True)
    p = doc.add_paragraph(style="摘要正文")
    p.add_run(strip_alnum_han_spaces(
        "面向多无人机在三维空间中的编队飞行与障碍规避问题，构建领航者—跟随者控制框架，并在一阶运动学模型上引入静态障碍物人工势场、无人机间分离项、局部极小检测与切向逃逸策略以及移动障碍物位置预测。系统采用统一的速度指令叠加方式协调目标跟踪、队形保持和安全避障，通过速度模长限幅和平滑更新保证轨迹连续。围绕编队增益、静态避障、局部极小和动态预测开展4组递进实验。结果表明，3组编队增益均能完成任务，增益提高可缩短完成时间并降低全过程编队误差；静态人工势场将障碍物碰撞次数由5次降为0，但增加了绕行距离和瞬时编队误差；切向逃逸策略使对称多障碍物场景下的批量成功率由10%提高到100%；移动障碍物实验的6个组合均无碰撞，但固定1s预测时域并未在所有速度和指标上占优。研究说明，人工势场适合作为轻量实时避障方法，其改进策略应结合场景几何关系和相对运动特征进行参数配置。"
    ))
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run("关键词：")
    set_run_font(r, size=12, bold=True)
    r = p.add_run("多无人机；三维编队；人工势场法；协同避障；局部极小；预测避障")
    set_run_font(r, size=12)
    doc.add_page_break()

    # 英文摘要
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Abstract")
    set_run_font(r, cn=FONT_CN, en=FONT_EN, size=16, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    text = (
        "A leader-follower framework is developed for three-dimensional formation flight and cooperative obstacle avoidance of multiple unmanned aerial vehicles. "
        "The controller combines formation tracking, artificial potential fields, inter-vehicle separation, local-minimum detection with tangential escape, and predicted-position avoidance for moving obstacles. "
        "Four progressive experiments evaluate formation gain, static obstacle avoidance, local-minimum mitigation, and dynamic obstacle prediction. "
        "All tested formation gains complete the nominal task. Static obstacle avoidance eliminates five collision events at the cost of longer routes and larger transient formation errors. "
        "Tangential escape improves the batch success rate from 10% to 100% in a symmetric multi-obstacle scenario. All dynamic-obstacle combinations remain collision-free, while a fixed 1s prediction horizon does not dominate every metric at every obstacle speed. "
        "The results indicate that artificial potential fields provide a lightweight real-time solution, but predictive and escape parameters must be matched to the encounter geometry."
    )
    r = p.add_run(text)
    set_run_font(r, cn=FONT_CN, en=FONT_EN, size=12)
    p = doc.add_paragraph()
    r = p.add_run("Key words: ")
    set_run_font(r, cn=FONT_CN, en=FONT_EN, size=12, bold=True)
    r = p.add_run("multi-UAV system; three-dimensional formation; artificial potential field; cooperative obstacle avoidance; local minimum; predictive avoidance")
    set_run_font(r, cn=FONT_CN, en=FONT_EN, size=12)
    doc.add_page_break()

    # 目录
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("目  录")
    set_run_font(r, size=16, bold=True)
    add_toc(doc)

    # 第一章
    add_heading(doc, "第一章 研究背景与意义", 1)
    add_heading(doc, "1.1 研究背景", 2)
    add_body(doc, "无人机集群通过多机协作扩展了单架无人机的覆盖范围和任务冗余，在巡检、测绘、搜索救援和复杂环境监测等任务中具有明显优势。与单机飞行相比，集群系统不仅需要控制个体到达目标，还要维持空间构型、限制机间距离，并在出现障碍物时协调安全性与队形完整性。上述目标相互耦合，使多无人机编队控制成为典型的多目标实时决策问题。")
    add_body(doc, "人工势场法将目标引力、障碍物排斥力和机间分离作用统一为局部控制量，计算结构简单，适合高频在线计算。Khatib提出的势场思想为实时避障提供了经典框架[1]；群集与一致性研究进一步说明，局部交互规则可以形成稳定的集体运动[2-4]。然而，标准人工势场存在局部极小和狭窄区域振荡等问题；面对移动障碍物时，仅依据当前位置计算排斥量还可能造成响应滞后。")
    add_heading(doc, "1.2 研究意义", 2)
    add_heading(doc, "1.2.1 理论意义", 3)
    add_body(doc, "本研究在同一控制框架中讨论编队保持、静态避障、机间防碰撞、局部极小逃逸和移动障碍物预测，能够直观展示不同控制项之间的竞争关系。通过递进实验比较基准方法与改进方法，可以区分算法有效性、适用条件和性能代价，避免只凭单条轨迹判断方法优劣。")
    add_heading(doc, "1.2.2 工程与教学意义", 3)
    add_body(doc, "课程项目采用Python、NumPy和Matplotlib实现，计算过程透明，参数可复现，轨迹、误差和安全距离均可直接可视化。该实现便于分析速度限幅、数值保护、随机种子控制和批量实验等工程细节，也为后续接入动力学模型、通信约束和物理仿真平台提供基础。")
    add_heading(doc, "1.3 研究目标与主要内容", 2)
    add_body(doc, "研究目标是在三维一阶运动学条件下，使1架领航者和4架跟随者形成楔形编队并到达目标区域，同时避免静态或移动球形障碍物以及无人机之间的几何碰撞。围绕该目标，项目完成以下工作：")
    add_bullets(doc, [
        "建立随领航者航向旋转的三维编队期望位置模型；",
        "设计目标吸引、编队保持、障碍物排斥和机间分离的复合速度控制器；",
        "利用连续停滞判据识别局部极小，并叠加确定性切向逃逸量；",
        "采用匀速外推获得移动障碍物预测位置，并与当前位置方法进行配对比较；",
        "构建统一指标体系和固定随机种子批量实验，分析安全性、效率和编队性能。",
    ])
    add_heading(doc, "1.4 论文结构", 2)
    add_body(doc, "全文共6章。第一章说明研究背景、意义与目标；第二章建立系统模型并定义评价指标；第三章给出编队控制、人工势场及两类改进策略；第四章说明实验设计和可复现性条件；第五章依据4组实验结果进行分析；第六章总结结论、局限与后续工作。")

    # 第二章
    add_heading(doc, "第二章 系统建模与问题描述", 1)
    add_heading(doc, "2.1 场景与基本假设", 2)
    add_body(doc, "系统包含5架无人机，其中0号无人机为领航者，1至4号无人机为跟随者。领航者初始位置为[0,0,12]m，目标位置为[50,0,12]m。跟随者在各自期望编队位置附近加入固定种子生成的初始扰动。障碍物采用球体描述，无人机采用等效球半径进行碰撞判断。")
    add_body(doc, "模型采用三维一阶运动学，不描述质量、惯性矩、螺旋桨推力、力矩或空气动力学。横滚角、俯仰角、偏航角和视觉旋翼转速只用于表现运动趋势，不参与轨迹状态更新。因此，本文中的势场“力”实质上是速度控制项，单位和物理含义与真实牛顿力不同。")
    add_heading(doc, "2.2 运动学模型", 2)
    add_heading(doc, "2.2.1 位置与速度更新", 3)
    add_body(doc, "设第i架无人机在离散时刻k的位置和速度分别为p_i(k)和v_i(k)，仿真步长为Δt，则位置按下式更新：")
    add_equation(doc, "pᵢ(k+1)=pᵢ(k)+vᵢ(k)Δt")
    add_body(doc, "控制器输出期望速度v_cmd,i(k)。为避免速度突变，实际使用的速度采用一阶平滑：")
    add_equation(doc, "vᵢ(k+1)=(1-α)vᵢ(k)+αv_cmd,ᵢ(k)")
    add_body(doc, "当速度模长超过v_max时，按向量模长整体缩放，从而保持运动方向不变。默认步长为0.05s，最大速度为3.0m/s。")
    add_heading(doc, "2.2.2 姿态与旋翼可视化", 3)
    add_body(doc, "显示姿态根据速度命令与当前速度之差估算近似加速度，并将横滚角和俯仰角限制在±25°。偏航角由水平速度方向确定，角度误差被映射到[-π,π]，以防止跨越±180°时出现不必要的大角度旋转。视觉旋翼相邻方向相反，但转速不会反作用于位置和速度。")
    add_heading(doc, "2.3 三维楔形编队", 2)
    add_body(doc, "编队在领航者机体系下定义。4架跟随者的相对偏移分别为[-5,3,2]m、[-5,-3,2]m、[-5,3,-2]m和[-5,-3,-2]m。设领航者航向角为ψ，绕z轴的旋转矩阵为R_z(ψ)，则第i架跟随者的期望位置为：")
    add_equation(doc, "p_d,ᵢ=p₀+R_z(ψ)δᵢ")
    add_body(doc, "该定义使楔形队列随领航者航向同步转动，而不是固定在世界坐标系中。跟随者既需要缩小位置误差，也需要跟踪领航者速度。")
    add_heading(doc, "2.4 任务目标与评价指标", 2)
    add_body(doc, "实验成功同时要求：在最大仿真时间内到达目标区域；不发生障碍物碰撞；不发生无人机碰撞；最终平均编队误差小于1m。编队误差定义为跟随者当前位置与期望编队位置的欧氏距离，并对4架跟随者取平均。")
    add_equation(doc, "E_f(k)=1/(N-1)·Σᵢ₌₁ᴺ⁻¹||pᵢ(k)-p_d,ᵢ(k)||")
    add_body(doc, "安全性由最小障碍物表面净距离和最小机间距离衡量；效率由完成时间和集群总航程衡量；动态场景还统计离开障碍物影响区域后的编队恢复时间。负的障碍物净距离表示几何重叠。")
    add_seq_caption(doc, "表", 1, "仿真系统的主要公共参数", "tbl1", before=True)
    add_three_line_table(
        doc,
        ["参数", "符号或取值", "说明"],
        [
            ["无人机数量", "5", "1架领航者、4架跟随者"],
            ["仿真步长", "0.05s", "离散状态更新时间"],
            ["最大速度", "3.0m/s", "按三维速度模长限幅"],
            ["目标位置", "[50,0,12]m", "领航者目标点"],
            ["最终误差阈值", "1m", "平均编队误差成功判据"],
            ["预测时域", "1.0s", "移动障碍物匀速外推时间"],
        ],
        [3.2, 4.2, 8.0],
        font_size=10,
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_body_with_refs(doc, ["主要公共参数汇总见表", ("ref", "tbl1", 1), "。所有对比方法在同组实验中使用相同初始条件，使差异主要来自控制策略。"])

    # 第三章
    add_heading(doc, "第三章 技术方案", 1)
    add_heading(doc, "3.1 系统总体架构", 2)
    add_heading(doc, "3.1.1 控制链与数据流", 3)
    add_body(doc, "系统从任务参数、障碍物状态和编队构型出发，计算目标吸引、编队修正、障碍物排斥和机间分离等控制项；随后完成局部极小判断、速度限幅和平滑更新，得到下一时刻的无人机状态，并在仿真结束后统计安全性和编队性能。")
    add_figure(doc, flow_path, width_cm=15.0, alt="多无人机编队与协同避障控制链流程图")
    add_seq_caption(doc, "图", 1, "多无人机编队与协同避障控制链", "fig1")
    add_body_with_refs(doc, ["如图", ("ref", "fig1", 1), "所示，轨迹状态在每个时间步反馈到控制器，形成闭环迭代。可视化模块读取状态历史，但不改变控制结果。"])
    add_heading(doc, "3.1.2 软件模块划分", 3)
    add_body(doc, "公共配置由config.py管理；models.py定义无人机和静态球形障碍物；controller.py实现编队控制、人工势场和显示姿态；simulator.py完成实验一。实验二至实验四分别由独立入口、仿真器和可视化模块组成，既复用公共控制逻辑，又隔离各实验的特有状态与指标。")
    add_heading(doc, "3.2 领航者—跟随者编队控制", 2)
    add_body(doc, "领航者由目标吸引项驱动，基础速度命令为k_goal(p_goal-p₀)。跟随者的基础速度命令由位置修正项和速度一致项构成：")
    add_equation(doc, "v_form,ᵢ=k_form(p_d,ᵢ-pᵢ)+k_vel(v₀-vᵢ)")
    add_body(doc, "位置修正项将跟随者拉向队形中的目标位置，速度一致项减小其与领航者之间的速度差。当k_form增大时，编队形成更快，但横向姿态调整也更积极，因此需要通过实验考察响应速度与平稳性之间的折中。")
    add_heading(doc, "3.3 静态障碍物人工势场", 2)
    add_heading(doc, "3.3.1 障碍物表面净距离", 3)
    add_body(doc, "设无人机中心位置为p，球形障碍物中心和半径分别为c和r_obs，无人机等效半径为r_uav，则表面净距离为：")
    add_equation(doc, "ρ=||p-c||-r_obs-r_uav")
    add_body(doc, "当ρ大于影响距离ρ₀时，障碍物不产生排斥；当ρ小于ρ₀时，排斥项随距离减小快速增大。程序对分母设置最小值并限制排斥项模长，以避免ρ接近0时发生数值爆炸。")
    add_heading(doc, "3.3.2 排斥项", 3)
    add_equation(doc, "F_obs=k_obs(1/ρ-1/ρ₀)(1/ρ²)·(p-c)/||p-c||，ρ<ρ₀")
    add_body(doc, "多个障碍物的排斥项直接累加。领航者将目标吸引、障碍排斥和机间分离相加；跟随者在此基础上还保留编队位置和速度一致控制。")
    add_heading(doc, "3.4 无人机间防碰撞", 2)
    add_body(doc, "当两架无人机之间的中心距离小于安全距离d_safe时，产生方向相反的分离项。该项与障碍物排斥具有相似的距离衰减形式，但作用对象是无人机对。完全重合时采用由编号决定的固定分离方向，避免随机数破坏可复现性。")
    add_equation(doc, "F_sep=k_sep(1/d-1/d_safe)(1/d²)·direction，d<d_safe")
    add_heading(doc, "3.5 局部极小检测与切向逃逸", 2)
    add_heading(doc, "3.5.1 连续停滞判据", 3)
    add_body(doc, "单个时间步的低速可能来自正常转向，不能直接判定为局部极小。系统保存最近20步的速度和任务距离，仅当无人机尚未接近目标、当前速度低于0.15m/s、附近存在障碍物、窗口内距离改善小于0.05m，并且上述条件连续满足指定步数时，才进入逃逸状态。")
    add_heading(doc, "3.5.2 切向逃逸方向", 3)
    add_body(doc, "设n为障碍物中心指向无人机的径向单位向量，则通过与竖直轴叉乘获得侧向单位向量t，并叠加大小为k_escape的切向速度项：")
    add_equation(doc, "t=normalize(n×[0,0,1])，F_escape=k_escape·t")
    add_body(doc, "当叉乘结果过小时改用[1,0,0]作为参考轴。绕行方向由无人机编号和领航者方向确定，使相同初始条件获得确定性结果。")
    add_heading(doc, "3.6 移动障碍物预测避障", 2)
    add_body(doc, "移动障碍物真实位置按恒速模型更新。当前位置方法直接使用c(k)计算排斥项；预测方法以固定时域T_pred外推未来位置：")
    add_equation(doc, "c_pred(k)=c(k)+T_pred·v_obs")
    add_body(doc, "预测位置只用于控制决策，障碍物真实轨迹和碰撞检测仍使用c(k)。该处理避免将预测误差掩盖为碰撞检测误差。固定时域能够提供提前量，但其效果与障碍物速度和相遇几何关系密切相关。")
    add_heading(doc, "3.7 数值稳定性与可复现性", 2)
    add_body(doc, "所有距离归一化均设置epsilon保护，速度按模长限幅，姿态角具有边界，批量运行每次重新创建无人机与障碍物。同一随机种子下的对比方法共享初始条件；关闭可视化参数不会改变轨迹。碰撞事件采用连续接触去重，同一对无人机的(i,j)和(j,i)只记录为一个事件。")

    # 第四章
    add_heading(doc, "第四章 实验设计", 1)
    add_heading(doc, "4.1 实验平台与评价方法", 2)
    add_body(doc, "仿真程序由Python实现，NumPy用于向量计算，Matplotlib用于三维轨迹、误差曲线和动画绘制，Pillow用于保存GIF。实验结果分别写入CSV、JSON和NPZ文件，图像仅用于展示，统计结论以结构化指标文件为依据。")
    add_heading(doc, "4.2 实验一：编队形成与保持", 2)
    add_body(doc, "在无障碍物环境中比较k_form=0.3、0.6和1.0。3组实验使用相同初始扰动，考察完成时间、全过程平均编队误差、最大误差、最终误差、总航程和姿态变化。该实验用于验证基础编队控制，并观察控制增益的影响。")
    add_figure(doc, ROOT / "results/experiment_1/k_form_0_6/trajectory_3d.png", width_cm=14.2, alt="编队增益0.6时的三维轨迹")
    add_seq_caption(doc, "图", 2, "编队增益为0.6时的三维飞行轨迹", "fig2")
    add_body_with_refs(doc, ["图", ("ref", "fig2", 2), "给出了代表性三维轨迹。领航者沿目标方向飞行，跟随者从带扰动的初始位置逐步收敛到楔形编队。"])
    add_heading(doc, "4.3 实验二：静态障碍物避障", 2)
    add_body(doc, "在[25,0,12]m处设置半径为4m的球形障碍物，比较无避障、仅障碍物人工势场、人工势场加机间防碰撞3种方法。重点统计障碍物碰撞次数、最小障碍物净距离、最小机间距离、完成时间、总航程和编队误差。")
    add_heading(doc, "4.4 实验三：局部极小改进", 2)
    add_body(doc, "在对称多障碍物场景中比较标准人工势场与加入切向逃逸的改进方法。每种方法运行10个固定随机种子，单次失败不终止批量实验。成功率为首要指标，同时统计成功样本完成时间、总航程、编队误差、局部极小检测次数和逃逸启用次数。")
    add_heading(doc, "4.5 实验四：移动障碍物预测", 2)
    add_body(doc, "移动障碍物从[25,-12,12]m出发，沿y轴正方向穿过集群路径。分别测试0.5m/s、1.0m/s和1.5m/s，每种速度比较当前位置与预测位置方法，并对每个组合运行5个固定随机种子，共30次运行。")

    # 第五章
    add_heading(doc, "第五章 结果与分析", 1)
    add_heading(doc, "5.1 编队增益对形成性能的影响", 2)
    exp1_rows = []
    for row in exp1:
        exp1_rows.append([
            row["k_form"],
            "是" if row["reached_goal_and_formation"] == "True" else "否",
            f'{float(row["completion_time_s"]):.2f}',
            f'{float(row["mean_formation_error_m"]):.3f}',
            f'{float(row["max_formation_error_m"]):.3f}',
            f'{float(row["final_formation_error_m"]):.3f}',
            f'{float(row["total_flight_distance_m"]):.3f}',
        ])
    add_seq_caption(doc, "表", 2, "不同编队增益下的主要结果", "tbl2", before=True)
    add_three_line_table(
        doc,
        ["k_form", "完成任务", "时间/s", "平均误差/m", "最大误差/m", "最终误差/m", "总航程/m"],
        exp1_rows,
        [1.8, 2.0, 2.0, 2.6, 2.6, 2.6, 2.6],
        font_size=8.8,
    )
    add_body_with_refs(doc, ["由表", ("ref", "tbl2", 2), "可见，3组参数均到达目标且最终平均编队误差小于1m。k_form由0.3增至1.0时，完成时间由26.75s缩短至18.35s，平均编队误差由6.055m降至2.662m。"])
    add_body(doc, "增益提高后跟随者对位置偏差的修正更快，但最大横滚角由1.882°上升至5.009°，姿态变化均方根也随之增加。因而较大增益改善收敛速度的同时，会提高控制动作强度。全过程平均误差包含初始编队形成阶段，不能当作稳态误差。")
    add_figure(doc, ROOT / "results/experiment_1/formation_error_comparison.png", width_cm=14.5, alt="不同编队增益的平均编队误差曲线")
    add_seq_caption(doc, "图", 3, "不同编队增益的平均编队误差对比", "fig3")
    add_body_with_refs(doc, ["图", ("ref", "fig3", 3), "显示3条误差曲线最终均收敛至1m以内，其中较大增益在早期阶段下降更快。"])

    add_heading(doc, "5.2 静态障碍物避障效果", 2)
    names2 = {"A_no_avoidance": "无避障", "B_obstacle_apf": "仅障碍APF", "C_apf_and_separation": "APF加机间分离"}
    exp2_rows = []
    for row in exp2:
        exp2_rows.append([
            names2[row["method"]],
            "是" if row["success"] == "True" else "否",
            f'{float(row["completion_time_s"]):.2f}',
            f'{float(row["minimum_obstacle_clearance_m"]):.3f}',
            f'{float(row["minimum_inter_uav_distance_m"]):.3f}',
            row["obstacle_collision_count"],
            f'{float(row["total_flight_distance_m"]):.3f}',
        ])
    add_seq_caption(doc, "表", 3, "静态障碍物场景的对比结果", "tbl3", before=True)
    add_three_line_table(
        doc,
        ["方法", "成功", "时间/s", "最小净距/m", "最小机间距/m", "碰撞次数", "总航程/m"],
        exp2_rows,
        [3.1, 1.5, 1.9, 2.4, 2.6, 2.0, 2.4],
        font_size=8.8,
    )
    add_body_with_refs(doc, ["表", ("ref", "tbl3", 3), "表明，无避障时虽然仍到达目标并恢复编队，但穿过障碍物并记录5次连续接触事件，最小表面净距离为-4.250m。启用人工势场后碰撞次数降至0，最小净距离提高到1.049m。"])
    add_body(doc, "人工势场使完成时间由20.65s增加到约37.6s，总航程由244.475m增加到约304.15m，说明安全绕行需要付出时间和路径代价。加入机间分离后，最小机间距离从2.948m提高到2.973m，增幅有限，因为该组初始条件本身未接近实际碰撞阈值。")
    add_figure(doc, ROOT / "results/experiment_2/trajectory_comparison_3d.png", width_cm=15.0, alt="静态障碍物场景三种方法的三维轨迹对比")
    add_seq_caption(doc, "图", 4, "静态障碍物场景的三维轨迹对比", "fig4")
    add_body_with_refs(doc, ["从图", ("ref", "fig4", 4), "可以观察到，无避障轨迹直接穿过球形区域；两种人工势场方法在障碍物前方产生侧向绕行，并在通过后恢复队形。"])

    add_heading(doc, "5.3 切向逃逸对局部极小的改善", 2)
    names3 = {"standard_apf": "标准APF", "improved_apf": "APF加切向逃逸"}
    exp3_rows = []
    for row in exp3:
        exp3_rows.append([
            names3[row["method"]],
            row["run_count"],
            f'{float(row["success_rate_percent"]):.0f}%',
            f'{float(row["average_completion_time_successful_s"]):.2f}',
            f'{float(row["average_total_flight_distance_m"]):.3f}',
            row["total_local_minimum_detections"],
            row["total_escape_activations"],
            row["total_collisions"],
        ])
    add_seq_caption(doc, "表", 4, "局部极小改进方法的批量结果", "tbl4", before=True)
    add_three_line_table(
        doc,
        ["方法", "次数", "成功率", "成功时间/s", "平均航程/m", "检测次数", "逃逸次数", "碰撞"],
        exp3_rows,
        [3.2, 1.2, 1.7, 2.2, 2.4, 1.8, 1.8, 1.4],
        font_size=8.8,
    )
    add_body_with_refs(doc, ["表", ("ref", "tbl4", 4), "显示，标准APF在10次运行中仅1次成功，成功率为10%；加入切向逃逸后10次全部成功，成功率达到100%，且两种方法均未发生碰撞。"])
    add_body(doc, "标准方法的平均航程和平均编队误差看似更小，原因不是其整体性能更优，而是9个失败样本较早停滞，累计运动距离和后续编队变化较少。改进方法共检测48次局部极小并启用48次逃逸，代价是更长的路径和更明显的姿态变化。该结果说明成功率应与效率指标联合解释。")
    add_figure(doc, ROOT / "results/experiment_3/success_rate.png", width_cm=12.5, alt="标准人工势场与切向逃逸方法的成功率对比")
    add_seq_caption(doc, "图", 5, "标准APF与切向逃逸方法的成功率", "fig5")
    add_body_with_refs(doc, ["图", ("ref", "fig5", 5), "直观展示了两种方法在对称多障碍物场景中的成功率差异。校验文件同时确认标准方法未使用逃逸力、20次运行均写入CSV且批量过程无异常。"])

    add_heading(doc, "5.4 移动障碍物预测方法的速度依赖", 2)
    names4 = {"current_position": "当前位置", "predicted_position": "预测位置"}
    exp4_rows = []
    for row in exp4:
        exp4_rows.append([
            f'{float(row["obstacle_speed_m_per_s"]):.1f}',
            names4[row["method"]],
            f'{float(row["success_rate_percent"]):.0f}%',
            f'{float(row["average_completion_time_successful_s"]):.2f}',
            f'{float(row["average_minimum_obstacle_clearance_m"]):.3f}',
            f'{float(row["average_max_formation_error_m"]):.3f}',
            f'{float(row["average_recovery_time_recovered_s"]):.2f}',
            f'{float(row["average_total_flight_distance_m"]):.3f}',
        ])
    add_seq_caption(doc, "表", 5, "移动障碍物场景的批量结果", "tbl5", before=True)
    add_three_line_table(
        doc,
        ["速度/(m/s)", "方法", "成功率", "时间/s", "净距/m", "最大误差/m", "恢复/s", "航程/m"],
        exp4_rows,
        [2.0, 2.2, 1.6, 1.8, 1.8, 2.2, 1.7, 2.2],
        font_size=8.4,
    )
    add_body_with_refs(doc, ["由表", ("ref", "tbl5", 5), "可知，6个速度—方法组合均达到100%成功率且碰撞次数为0。预测策略的收益随速度变化，并非所有指标单调改善。"])
    add_body(doc, "障碍物速度为0.5m/s时，预测位置将平均最小净距离由1.696m提高到1.899m；速度为1.0m/s时，预测方法的平均最小净距离反而降至0.955m，最大编队误差增至8.935m；速度为1.5m/s时，预测方法虽未获得更大的净距离，却将完成时间由23.00s缩短至21.65s，将恢复时间由10.70s缩短至10.10s，并减少了总航程和最大编队误差。")
    add_body(doc, "上述差异源于固定1s预测时域改变了排斥作用的触发位置。预测位置相对真实位置的位移随速度增大，可能使无人机提前绕行，也可能改变随后与障碍物的交汇关系。故动态避障不宜仅比较单一最小净距离，还应综合任务时间、编队变形和恢复过程。")
    add_figure(doc, ROOT / "results/experiment_4/minimum_clearance_comparison.png", width_cm=14.5, alt="不同移动障碍物速度下的最小净距离对比")
    add_seq_caption(doc, "图", 6, "不同障碍物速度下的最小净距离对比", "fig6")
    add_body_with_refs(doc, ["图", ("ref", "fig6", 6), "进一步表明固定预测时域具有速度依赖。30次运行均完成编队恢复，且校验结果确认碰撞检测始终使用真实障碍物位置。"])

    add_heading(doc, "5.5 综合讨论", 2)
    add_heading(doc, "5.5.1 安全性与编队性能的权衡", 3)
    add_body(doc, "实验二至实验四共同说明，避障阶段允许编队暂时变形是合理的控制选择。障碍物排斥与编队保持的方向可能冲突，若强制保持紧密队形，边缘无人机可能失去足够安全裕度；若排斥项过强，则会产生更长路径和更大瞬时误差。控制器的目标应是优先避免碰撞，在脱离风险区后恢复编队。")
    add_heading(doc, "5.5.2 改进策略的适用条件", 3)
    add_body(doc, "切向逃逸直接针对标准人工势场在对称障碍物附近的停滞问题，因此成功率提升明显；位置预测针对的是运动障碍物的时间提前量，其效果同时受速度、横穿时刻和预测时域影响。两种改进的作用机制不同，不能用同一指标判断。")
    add_heading(doc, "5.5.3 结果可信度", 3)
    add_body(doc, "实验通过固定随机种子、配对初始条件、结构化结果文件和独立验证项提高可复现性。实验三验证了切向向量与径向向量正交，实验四验证了预测过程不修改真实障碍物轨迹；两组批量实验的仿真异常数均为0。现有结论适用于所设运动学模型和参数范围，不应外推为真实飞行器的绝对性能。")

    # 第六章
    add_heading(doc, "第六章 结论与展望", 1)
    add_heading(doc, "6.1 主要结论", 2)
    add_body(doc, "本文完成了5架无人机三维楔形编队与协同避障仿真。领航者—跟随者控制器能够在无障碍物环境中形成并保持编队，3组编队增益均满足最终误差要求；提高编队增益能够加快收敛，但会增强姿态变化。")
    add_body(doc, "静态障碍物人工势场将碰撞次数由5次降为0，证明其能够实现实时绕障，同时也增加任务时间、航程和瞬时编队误差。机间分离项在本场景中略微提高最小机间距离，其效果受初始构型和安全阈值影响。")
    add_body(doc, "针对局部极小，连续停滞检测与切向逃逸将批量成功率由10%提高到100%，但增加了飞行距离和姿态变化。针对移动障碍物，固定1s预测时域在部分速度下改善提前性和恢复过程，但不能保证所有指标均优于当前位置方法。")
    add_heading(doc, "6.2 研究局限", 2)
    add_body(doc, "当前模型采用一阶运动学，未描述推力、力矩、惯性、空气阻力和电机响应，因此不能据此评价真实能耗或飞行稳定裕度。障碍物为规则球体，通信时延、定位噪声、风扰动和传感器误差均未纳入。批量样本数量和参数范围有限，统计结果主要用于课程实验中的方法比较。")
    add_heading(doc, "6.3 后续工作", 2)
    add_body(doc, "后续可将控制器接入六自由度四旋翼动力学模型，并引入传感器噪声、通信丢包和时延；在复杂障碍物环境中研究自适应势场参数、基于相对速度的碰撞时间预测和多目标任务分配；进一步在Webots或Gazebo等物理仿真平台中进行软硬件在环验证。")

    # 参考文献
    add_heading(doc, "参考文献", 1)
    references = [
        "[1] Khatib O. Real-time obstacle avoidance for manipulators and mobile robots[J]. The International Journal of Robotics Research,1986,5(1):90-98.",
        "[2] Reynolds C W. Flocks, herds and schools: A distributed behavioral model[C]//Proceedings of SIGGRAPH.1987:25-34.",
        "[3] Olfati-Saber R. Flocking for multi-agent dynamic systems: Algorithms and theory[J]. IEEE Transactions on Automatic Control,2006,51(3):401-420.",
        "[4] Ren W, Beard R W. Distributed Consensus in Multi-vehicle Cooperative Control[M]. London: Springer,2008.",
        "[5] Beard R W, McLain T W. Small Unmanned Aircraft: Theory and Practice[M]. Princeton: Princeton University Press,2012.",
        "[6] LaValle S M. Planning Algorithms[M]. Cambridge: Cambridge University Press,2006.",
    ]
    for ref in references:
        p = doc.add_paragraph(style="正文")
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.first_line_indent = Pt(-24)
        r = p.add_run(ref)
        set_run_font(r, size=10.5)

    for section in doc.sections:
        add_footer(section)

    disable_auto_spacing_everywhere(doc)

    doc.save(OUTPUT)
    return OUTPUT


def verify_ooxml(path):
    required = {
        "TOC": False,
        "SEQ 图": False,
        "SEQ 表": False,
        "REF fig1": False,
        "REF tbl1": False,
        "updateFields": False,
    }
    with zipfile.ZipFile(path) as zf:
        document_xml = zf.read("word/document.xml").decode("utf-8")
        settings_xml = zf.read("word/settings.xml").decode("utf-8")
        for key in list(required):
            hay = settings_xml if key == "updateFields" else document_xml
            required[key] = key in hay
        styles_xml = zf.read("word/styles.xml").decode("utf-8")
        if "Heading1" not in styles_xml or "Heading2" not in styles_xml or "Heading3" not in styles_xml:
            raise RuntimeError("标题样式缺失")
    missing = [k for k, v in required.items() if not v]
    if missing:
        raise RuntimeError(f"字段或设置缺失:{missing}")
    return required


if __name__ == "__main__":
    out = build_document()
    checks = verify_ooxml(out)
    print(f"OUTPUT={out}")
    print(f"SIZE={out.stat().st_size}")
    print("CHECKS=" + json.dumps(checks, ensure_ascii=False))
